"""Build the PSH technical report as a styled DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "PSH_Planner_技术报告.docx"
ASSETS = ROOT / "docs" / "assets"

NAVY = "16324F"
BLUE = "2F6B9A"
TEAL = "2A9D8F"
RED = "C94C4C"
INK = "263238"
MUTED = "607080"
LIGHT = "EAF2F8"
GRID = "D9E1E8"


def set_font(run, name="Microsoft YaHei", size=10.5, bold=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size=8):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PSH Planner  |  ")
    set_font(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_para(doc, text="", bold_lead=None, align=None, space_after=6, first_indent=True):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, size=16 if level == 1 else 12.5, bold=True, color="000000")
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(item)
        set_font(run)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, size=9, color=MUTED)
    return p


def add_figure(doc, path, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    picture = p.add_run().add_picture(str(path), width=Cm(width_cm))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    add_caption(doc, caption)


def add_table(doc, headers, rows, widths=None, highlight_last=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_props = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_props.append(repeat)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Cm(widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(header))
        set_font(run, size=9, bold=True, color="FFFFFF")
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        row_props = table.rows[-1]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_props.append(cant_split)
        fill = "FFFFFF" if row_idx % 2 == 0 else "F3F7FA"
        if highlight_last and row_idx == len(rows) - 1:
            fill = "DDF2EF"
        for i, value in enumerate(values):
            cell = cells[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            if widths:
                cell.width = Cm(widths[i])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_font(run, size=8.8, bold=(highlight_last and row_idx == len(rows) - 1))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    title = doc.styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    for name, size, before, after in (("Heading 1", 16, 14, 7), ("Heading 2", 12.5, 10, 4)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    configure_styles(doc)
    add_page_number(section.footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.style = doc.styles["Title"]
    run = p.add_run("预测式安全混合动态路径规划技术报告")
    set_font(run, size=24, bold=True, color="000000")
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(12)
    p2.paragraph_format.space_after = Pt(32)
    run = p2.add_run("局部观测动态障碍导航的架构设计与消融实验")
    set_font(run, size=15, color=MUTED)
    add_figure(doc, ASSETS / "architecture.png", 15.8, "图 1  PSH 系统结构与控制信息流")
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_before = Pt(24)
    run = p3.add_run("项目技术报告  版本 1.0  2026 年 9 月")
    set_font(run, size=10, color=MUTED)
    doc.add_page_break()

    add_heading(doc, "摘要", 1)
    add_para(doc, "本项目研究局部观测条件下的动态障碍路径规划。原型最初包含 A*、D* Lite、Q-learning、DQN、PPO 与 MCTS 等算法，但各模块主要用于独立演示，动态场景中的观测条件、随机轨迹和评价流程并不统一，难以回答清晰的研究问题。升级后的 Predictive Safe Hybrid Planner（PSH）将问题收敛为一套可验证架构：D* Lite 维护静态地图上的全局走廊；轻量障碍跟踪器根据连续局部观测估计速度并生成短时概率占据图；滚动时域 Beam Search 在目标进展、走廊偏离、动作代价和碰撞风险之间优化；独立安全盾在动作输出前执行最后的风险检查。")
    add_para(doc, "在 30 个未见随机种子上，完整 PSH 的成功率为 86.7%，碰撞率为 13.3%，平均单步规划延迟为 4.50 ms。相同条件下，在线 D* Lite 的成功率为 73.3%，碰撞率为 23.3%；在线重规划 A* 的成功率为 63.3%，碰撞率为 33.3%。去掉预测或安全盾后，成功率分别降至 70.0% 和 66.7%。这些结果支持项目的核心假设：短时预测与独立安全约束的协同能在毫秒级预算内改善动态导航可靠性。结论仅适用于当前离散栅格仿真，不等同于实体机器人验证。")
    add_heading(doc, "关键词", 2)
    add_para(doc, "动态路径规划  局部观测  D* Lite  概率占据预测  滚动时域规划  安全盾", first_indent=False)

    add_heading(doc, "1 问题定义与研究假设", 1)
    add_para(doc, "环境由静态栅格、移动障碍、起点和目标点组成。机器人状态包含行、列和离散朝向，动作集合为前进、后退、左转和右转。机器人只能观察以自身为中心、半径为 5 个栅格的移动障碍，静态地图已知。每执行一个动作，移动障碍同步更新位置，因此只依据当前占据状态做反应会遭遇典型的时序碰撞：目标栅格在决策时为空，但障碍在本周期结束时进入该位置。")
    add_para(doc, "研究假设是：在相同局部观测、动作上限和动态轨迹下，将短时运动预测、全局路线引导与独立动作级安全约束组合起来，可以降低碰撞率并提高到达率，同时将单步决策时间控制在实时软件可接受的毫秒量级。该假设通过在线基线和组件消融检验。")
    add_heading(doc, "2 原型审计与升级目标", 1)
    add_para(doc, "原型具有完整的地图生成、局部观测、传统搜索、强化学习和交互演示模块，为升级提供了可复用基础。审计同时发现几项会削弱实验可信度的问题。")
    add_table(doc,
        ["原型问题", "影响", "升级处理"],
        [
            ["动态环境 reset 不恢复障碍初态", "算法可能面对不同轨迹", "保存世界快照并重置局部随机数生成器"],
            ["随机游走使用全局 random", "结果无法严格复现", "每个环境使用独立带种子 RNG"],
            ["动态 A* 未纳入观测到的移动障碍", "基线信息条件不公平", "所有在线规划器共享同一局部观测流"],
            ["D* Lite 每步重新实例化或由 A* 代执行", "无法体现增量规划", "修复队列和边界并保持跨步状态"],
            ["报告以算法横向比较为主", "缺乏可检验方法主线", "建立 PSH 架构并进行模块消融"],
        ], widths=[4.2, 4.2, 7.2])

    add_heading(doc, "3 系统架构", 1)
    add_para(doc, "PSH 采用全局与局部分层设计。全局层只处理静态几何和长期可达性，局部层在每个控制周期处理移动障碍风险。安全盾独立于路径优化器，使错误预测或代价权重失配不会直接变成未经检查的控制命令。")
    add_bullets(doc, [
        "全局走廊提供从当前位置到目标的拓扑方向，避免局部搜索长期绕行。",
        "概率占据图把障碍历史转化为时间相关风险，而不是把移动障碍永久标成墙。",
        "滚动时域规划只执行最优序列的首个动作，新观测到达后立即重算。",
        "安全盾根据一步风险否决动作，并从可行动作中选择最低风险回退。",
        "DQN 可作为动作先验参与排序，但不能绕过确定性搜索与安全盾。",
    ])

    add_heading(doc, "4 方法设计", 1)
    add_heading(doc, "4.1 增量全局走廊", 2)
    add_para(doc, "D* Lite 从目标反向维护 g 值和一步前瞻 rhs 值。当机器人起点前移时，算法通过 km 修正启发式键值，仅修复受影响节点。实现中修正了越界邻居、重复优先队列条目失效和扩展计数，使规划器能够跨控制周期保留状态。全局走廊不直接决定动作，而是进入局部代价函数。")
    add_heading(doc, "4.2 障碍跟踪与概率占据", 2)
    add_para(doc, "跟踪器使用最近邻关联连续观测。对匹配障碍，速度由相邻时刻位置差估计，并限制在每轴一个栅格。未来 H 步采用常速度外推，中心格获得主要概率质量，相邻格承担随时间增长的不确定性。短暂漏检允许轨迹保留两个周期，防止障碍刚离开视野就被立刻遗忘。该模型有意保持轻量，以便每个控制周期重新生成风险图。")
    add_heading(doc, "4.3 滚动时域局部搜索", 2)
    add_para(doc, "局部规划器用 Beam Search 展开长度为 7 的动作序列，每层保留 72 个低代价且状态不同的候选。候选累计代价如下，其中 d_goal 为目标曼哈顿距离，d_corridor 为到全局走廊的距离，p_collision 为对应时刻的预测占据概率，n_visit 为历史访问次数，prior 为可选学习先验。")
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eq.paragraph_format.space_before = Pt(6)
    eq.paragraph_format.space_after = Pt(9)
    run = eq.add_run("J = Σₜ [c(aₜ) + w_g d_goal + w_c d_corridor + w_r p_collision + w_v n_visit - w_p prior]")
    set_font(run, name="Cambria Math", size=11.5, color="000000")
    add_para(doc, "重复访问惩罚用于抑制动态障碍附近的局部振荡。动作代价对转向和后退施加额外开销。预测风险权重大于路线偏离代价，使规划器可以临时离开全局走廊；障碍通过后，全局走廊重新主导向目标的进展。")
    add_heading(doc, "4.4 动作级安全盾", 2)
    add_para(doc, "安全盾使用单步预测风险和当前可见障碍检查候选动作。当风险达到阈值 0.45 时，盾牌重新评估所有动作，选择风险最低的可行动作。盾牌与局部搜索分开实现，便于记录干预次数，并通过消融区分预测优化与最终动作约束的贡献。")

    add_heading(doc, "5 工程实现与可复现性", 1)
    add_para(doc, "新代码位于 psh_planner 包中，预测、局部搜索、安全盾、混合控制和基线各自独立。benchmark_psh.py 负责构造场景、重置环境、同步观测、计时和聚合指标。每个规划器在同一 seed 下重新创建等价环境；地图、移动障碍初始位置及随机游走均由该 seed 控制。逐回合 CSV 保留原始结果，聚合 JSON 可由脚本重新生成。")
    add_table(doc,
        ["模块", "职责", "主要接口"],
        [
            ["prediction.py", "关联观测并输出 H 步风险图", "update  predict"],
            ["local_planner.py", "滚动时域 Beam Search", "plan"],
            ["safety.py", "动作否决与最低风险回退", "filter"],
            ["hybrid.py", "维护全局走廊与调度各模块", "reset  act"],
            ["baselines.py", "公平的在线 A* 与 D* Lite", "reset  act"],
            ["benchmark_psh.py", "多种子实验 指标与消融", "run_episode  benchmark"],
        ], widths=[3.6, 7.4, 4.6])

    doc.add_page_break()
    add_heading(doc, "6 实验设计", 1)
    add_heading(doc, "6.1 场景与控制变量", 2)
    add_table(doc,
        ["项目", "设置"],
        [
            ["测试种子", "20 至 49 共 30 个未见种子"],
            ["地图", "24 x 32 栅格  34 个静态短障碍段"],
            ["动态障碍", "4 个  水平 垂直 随机 水平"],
            ["局部观测", "半径 5 个栅格"],
            ["动作上限", "每回合 350 步"],
            ["PSH 搜索", "预测时域 7  Beam 宽度 72"],
            ["评价指标", "成功 碰撞 步数 路径长度 转向 最小间距 平均与 P95 延迟"],
        ], widths=[4.2, 11.4])
    add_para(doc, "所有方法只能使用同一局部移动障碍观测。在线 A* 每步将当前可见动态障碍临时写入规划图并重新搜索；在线 D* Lite 对可见障碍执行增量边代价更新。PSH 无预测消融只使用当前占据图，PSH 无安全盾保留预测但直接输出局部搜索动作。")
    doc.add_page_break()
    add_heading(doc, "6.2 定量结果", 2)
    add_figure(doc, ASSETS / "benchmark.png", 13.5, "图 2  基线与消融实验的成功率和碰撞率")
    add_table(doc,
        ["方法", "成功率", "碰撞率", "平均步数", "最小间距", "平均延迟", "P95 延迟"],
        [
            ["在线重规划 A*", "63.3%", "33.3%", "51.77", "1.93", "0.29 ms", "0.73 ms"],
            ["在线 D* Lite", "73.3%", "23.3%", "61.40", "2.09", "0.16 ms", "0.39 ms"],
            ["PSH 无预测", "70.0%", "30.0%", "40.90", "1.92", "4.62 ms", "5.99 ms"],
            ["PSH 无安全盾", "66.7%", "33.3%", "39.63", "1.97", "4.65 ms", "6.01 ms"],
            ["PSH 完整系统", "86.7%", "13.3%", "46.90", "2.10", "4.50 ms", "6.06 ms"],
        ], widths=[3.5, 1.8, 1.8, 2.0, 2.0, 2.2, 2.2], highlight_last=True)
    add_para(doc, "完整 PSH 相对在线 D* Lite 将成功率提高 13.4 个百分点，将碰撞率降低 10.0 个百分点。相对在线重规划 A*，成功率提高 23.4 个百分点，碰撞率降低 20.0 个百分点。PSH 的平均最小间距为 2.10 格，是五种设置中最高值。代价是单步规划延迟增加到约 4.5 ms，但 P95 仍为 6.06 ms。")
    add_heading(doc, "6.3 消融解释", 2)
    add_para(doc, "去掉预测后，成功率从 86.7% 降至 70.0%，碰撞率从 13.3% 上升到 30.0%。这说明只对当前占据状态反应不足以处理障碍在动作结束时进入机器人位置的时序碰撞。去掉安全盾后，成功率降至 66.7%，碰撞率升至 33.3%。局部优化会在路线进展和预测风险之间权衡，独立盾牌则提供硬阈值检查；两者承担不同职责。完整系统平均每回合触发 0.9 次安全盾干预，少量干预即可改变一部分关键回合结果。")
    add_para(doc, "本实验只报告 30 个固定种子的描述性统计，没有声称统计显著性。完整逐回合数据已保留，后续可扩大样本并采用配对置信区间或 Bootstrap 检验。")

    add_heading(doc, "7 典型场景", 1)
    add_figure(doc, ASSETS / "scenario_trace.png", 14.5, "图 3  Seed 24 中完整 PSH 的执行轨迹  红点为动态障碍初始位置")
    add_para(doc, "图 3 展示 seed 24 的一次成功执行。机器人从左上区域出发，先沿静态通道下行，再在中部转向，最后沿底部走廊到达右下目标。该回合共执行 59 个动作，最小动态障碍间距为 1.0 格。路径并非静态最短路的简单回放，而是在局部风险和全局走廊之间反复重算得到。")

    add_heading(doc, "8 创新价值与适用边界", 1)
    add_para(doc, "项目的创新价值主要在架构和验证方法，而不是提出新的最短路定理。第一，系统把动态障碍从二值墙提升为带时间索引的概率风险，使搜索能够区分障碍现在的位置和未来可能的位置。第二，学习策略被限制为可替换的动作先验，最终控制仍经过显式代价和安全约束，降低了端到端策略难以解释的问题。第三，统一 Planner 接口、确定性环境重置和逐回合原始数据让基线与消融能够在同一轨迹上比较。")
    add_para(doc, "当前结论受离散模型限制。障碍只有格点位置，没有速度噪声、加速度或形状；机器人也没有半径和动力学约束。匿名最近邻关联在障碍密集交叉时可能交换轨迹。安全盾只检查一步风险，不能证明全局可达或连续时间无碰撞。因此该系统适合作为算法原型和软件架构验证，不应直接宣称已经完成实体机器人部署。")

    add_heading(doc, "9 后续工作", 1)
    add_bullets(doc, [
        "将栅格风险图替换为带协方差的连续状态估计，并加入时间戳和漏帧处理。",
        "把局部 Beam Search 升级为满足速度和加速度约束的轨迹优化器。",
        "加入机器人圆形或多边形足迹，进行连续时间碰撞检测与安全距离约束。",
        "在多机器人场景引入意图预测或互惠避碰，处理相互影响的运动决策。",
        "扩大随机种子和场景难度，报告配对置信区间、失败类型和参数敏感性。",
    ])

    add_heading(doc, "10 结论", 1)
    add_para(doc, "PSH 将早期多算法原型重构为围绕动态导航安全性的研究项目。实现不再依赖单一算法胜负，而是用全局走廊、概率预测、滚动时域规划和动作级安全盾形成闭环。30 个未见种子的结果及消融表明，完整系统在当前仿真设置下同时提高到达率、降低碰撞率，并维持毫秒级在线计算。实验也暴露了剩余失败和离散模型边界，为连续轨迹规划和实体平台迁移提供了明确方向。")

    add_heading(doc, "参考文献", 1)
    refs = [
        "[1] Koenig S and Likhachev M. D Star Lite. Proceedings of the AAAI Conference on Artificial Intelligence. 2002.",
        "[2] Fox D Burgard W and Thrun S. The Dynamic Window Approach to Collision Avoidance. IEEE Robotics and Automation Magazine. 1997. 4(1): 23-33. DOI 10.1109/100.580977.",
        "[3] van den Berg J Guy S J Lin M and Manocha D. Reciprocal n body Collision Avoidance. Robotics Research. Springer Tracts in Advanced Robotics 70. 2011. 3-19.",
        "[4] LaValle S M. Planning Algorithms. Cambridge University Press. 2006.",
        "[5] Sutton R S and Barto A G. Reinforcement Learning An Introduction. Second Edition. MIT Press. 2018.",
    ]
    for ref in refs:
        add_para(doc, ref, first_indent=False, space_after=4)

    add_heading(doc, "附录 复现实验", 1)
    add_para(doc, "安装依赖并运行测试：", first_indent=False)
    for command in (
        "python -m pip install -r requirements.txt",
        "python -m unittest discover -s tests -v",
        "python benchmark_psh.py --quick --seeds 20 21 22",
    ):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(command)
        set_font(run, name="Consolas", size=9, color=NAVY)
    add_para(doc, "正式结果位于 results/psh_benchmark/episodes.csv 和 summary.json。测试覆盖动态环境重置、D* Lite 障碍更新、速度预测、安全盾拒绝和静态障碍仿真。", first_indent=False)

    props = doc.core_properties
    props.title = "预测式安全混合动态路径规划技术报告"
    props.subject = "局部观测动态障碍导航的架构设计与消融实验"
    props.keywords = "PSH, dynamic path planning, D Star Lite, safety shield"
    props.author = "PSH Planner Project"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
